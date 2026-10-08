"""
ATS Resume Checker
------------------
Upload a resume (PDF / DOCX / TXT), optionally paste a job description,
and get an ATS score plus concrete improvement suggestions from Gemini.
"""

import io
import json
import os
import re
import time

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
DEFAULT_MODEL = "gemini-3.8-flash"
# Tried in order when the chosen model is unavailable, overloaded or out of quota
FALLBACK_MODELS = ["gemini-3.7-flash", "gemini-3.5-flash-lite", "gemini-flash-latest"]
RETRIES_PER_MODEL = 3      # attempts per model for temporary errors (503 etc.)
BACKOFF_SECONDS = 2        # waits 2s, 4s between attempts
MAX_FILE_MB = 5
MAX_CHARS = 20000  # keep the prompt small and cheap

st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")


# ----------------------------------------------------------------------------
# Helpers: API key, text extraction, JSON parsing
# ----------------------------------------------------------------------------
def get_api_key() -> str | None:
    """Look for the key in Streamlit secrets, then env vars."""
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass  # no secrets.toml file present
    return os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")


def get_model_name() -> str:
    try:
        if "GEMINI_MODEL" in st.secrets:
            return st.secrets["GEMINI_MODEL"]
    except Exception:
        pass
    return os.getenv("GEMINI_MODEL", DEFAULT_MODEL)


def extract_text(filename: str, data: bytes) -> str:
    """Return plain text from a PDF, DOCX or TXT file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        pages = [(p.extract_text() or "") for p in reader.pages]
        return "\n".join(pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        # resumes often keep content inside tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Please upload PDF, DOCX or TXT.")


def parse_json(raw: str) -> dict:
    """Parse the model output, tolerating ```json fences or extra text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise


def clamp_score(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def as_list(value) -> list:
    if isinstance(value, list):
        return [str(v) for v in value if str(v).strip()]
    if isinstance(value, str) and value.strip():
        return [value]
    return []


# ----------------------------------------------------------------------------
# Gemini call
# ----------------------------------------------------------------------------
PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and career coach.
Analyse the resume below{jd_clause} and respond with ONLY a JSON object, no markdown.

Scoring rubric (total 100):
- keywords_relevance: 30  (role-relevant skills/keywords{jd_hint})
- formatting_structure: 20 (clear sections, parseable layout, consistent dates)
- experience_impact: 20   (action verbs, quantified achievements)
- skills_education: 15    (skills clearly listed, education complete)
- contact_completeness: 10 (name, email, phone, LinkedIn/GitHub)
- length_readability: 5   (concise, no typos, good length)

Be honest and strict. Do not invent facts that are not in the resume.

JSON schema:
{{
  "overall_score": <int 0-100>,
  "summary": "<2-3 sentence overview>",
  "section_scores": {{
    "keywords_relevance": <int 0-30>,
    "formatting_structure": <int 0-20>,
    "experience_impact": <int 0-20>,
    "skills_education": <int 0-15>,
    "contact_completeness": <int 0-10>,
    "length_readability": <int 0-5>
  }},
  "strengths": ["..."],
  "weaknesses": ["..."],
  "missing_keywords": ["..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "section": "...", "issue": "...", "suggestion": "..."}}
  ],
  "rewrite_examples": [
    {{"before": "<weak line from the resume>", "after": "<stronger version>"}}
  ]
}}

{jd_block}RESUME:
\"\"\"
{resume}
\"\"\"
"""


def error_kind(exc: Exception) -> str:
    """Classify a Gemini error: 'retry' (temporary), 'switch' (try another model) or 'fatal'."""
    msg = str(exc)
    low = msg.lower()
    if any(k in msg for k in ("503", "500", "504", "UNAVAILABLE", "INTERNAL", "DEADLINE_EXCEEDED")) \
            or any(k in low for k in ("overloaded", "high demand", "timed out", "timeout", "connection")):
        return "retry"
    if any(k in msg for k in ("404", "NOT_FOUND", "429", "RESOURCE_EXHAUSTED")) \
            or "no longer available" in low:
        return "switch"
    return "fatal"  # bad API key, invalid request, etc. Retrying will not help


def generate_with_fallback(client, model: str, prompt: str, config):
    """Call Gemini. Retry temporary errors with backoff, then fall back to other models."""
    last_error = None
    for candidate in [model] + [m for m in FALLBACK_MODELS if m != model]:
        for attempt in range(RETRIES_PER_MODEL):
            try:
                return client.models.generate_content(model=candidate, contents=prompt, config=config)
            except Exception as e:
                kind = error_kind(e)
                if kind == "fatal":
                    raise
                last_error = e
                if kind == "switch":
                    break  # no point retrying this model
                if attempt < RETRIES_PER_MODEL - 1:
                    time.sleep(BACKOFF_SECONDS * (2 ** attempt))
    raise RuntimeError(
        "Gemini is overloaded or unavailable right now (all models tried). "
        f"Please try again in a minute. Last error: {last_error}"
    )


def analyse_resume(resume_text: str, job_desc: str, api_key: str, model: str) -> dict:
    jd_clause = " against the job description" if job_desc else ""
    jd_hint = " matched to the job description" if job_desc else ""
    jd_block = f'JOB DESCRIPTION:\n"""\n{job_desc[:MAX_CHARS]}\n"""\n\n' if job_desc else ""
    prompt = PROMPT.format(
        jd_clause=jd_clause,
        jd_hint=jd_hint,
        jd_block=jd_block,
        resume=resume_text[:MAX_CHARS],
    )

    client = genai.Client(api_key=api_key)
    # Gemini 3 models work best with the default temperature, so we don't override it.
    config = types.GenerateContentConfig(response_mime_type="application/json")

    response = generate_with_fallback(client, model, prompt, config)
    data = parse_json(response.text)

    # Normalise so the UI never crashes on odd model output
    data["overall_score"] = clamp_score(data.get("overall_score"))
    data["summary"] = str(data.get("summary", ""))
    data["section_scores"] = data.get("section_scores") or {}
    for key in ("strengths", "weaknesses", "missing_keywords"):
        data[key] = as_list(data.get(key))
    data["improvements"] = [i for i in (data.get("improvements") or []) if isinstance(i, dict)]
    data["rewrite_examples"] = [r for r in (data.get("rewrite_examples") or []) if isinstance(r, dict)]
    return data


# ----------------------------------------------------------------------------
# UI
# ----------------------------------------------------------------------------
SECTION_MAX = {
    "keywords_relevance": 30,
    "formatting_structure": 20,
    "experience_impact": 20,
    "skills_education": 15,
    "contact_completeness": 10,
    "length_readability": 5,
}
PRIORITY_ICON = {"high": "🔴", "medium": "🟠", "low": "🟢"}


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 60:
        return "Good, needs polish"
    if score >= 40:
        return "Needs work"
    return "Poor"


def render_results(result: dict) -> None:
    score = result["overall_score"]
    left, right = st.columns([1, 2])
    with left:
        st.metric("ATS Score", f"{score}/100", score_label(score), delta_color="off")
        st.progress(score / 100)
    with right:
        st.subheader("Summary")
        st.write(result["summary"] or "No summary returned.")

    st.divider()
    st.subheader("Score breakdown")
    cols = st.columns(3)
    for i, (key, maximum) in enumerate(SECTION_MAX.items()):
        try:
            value = max(0, min(maximum, int(float(result["section_scores"].get(key, 0)))))
        except (TypeError, ValueError):
            value = 0
        with cols[i % 3]:
            st.write(f"**{key.replace('_', ' ').title()}**  ({value}/{maximum})")
            st.progress(value / maximum)

    st.divider()
    c1, c2 = st.columns(2)
    with c1:
        st.subheader("✅ Strengths")
        for s in result["strengths"] or ["None listed"]:
            st.markdown(f"- {s}")
    with c2:
        st.subheader("⚠️ Weaknesses")
        for w in result["weaknesses"] or ["None listed"]:
            st.markdown(f"- {w}")

    if result["missing_keywords"]:
        st.subheader("🔑 Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    st.subheader("🛠️ Improvements")
    order = {"high": 0, "medium": 1, "low": 2}
    items = sorted(
        result["improvements"],
        key=lambda i: order.get(str(i.get("priority", "")).lower(), 3),
    )
    if not items:
        st.info("No specific improvements returned.")
    for item in items:
        icon = PRIORITY_ICON.get(str(item.get("priority", "")).lower(), "⚪")
        title = f"{icon} {item.get('priority', '')} · {item.get('section', 'General')}"
        with st.expander(title):
            st.markdown(f"**Issue:** {item.get('issue', '')}")
            st.markdown(f"**Fix:** {item.get('suggestion', '')}")

    if result["rewrite_examples"]:
        st.subheader("✍️ Rewrite examples")
        for ex in result["rewrite_examples"]:
            st.markdown(f"**Before:** {ex.get('before', '')}")
            st.markdown(f"**After:** {ex.get('after', '')}")
            st.write("")

    st.download_button(
        "⬇️ Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume and get an ATS score with practical improvements, powered by Gemini.")

    api_key = get_api_key()
    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input("Gemini API key", type="password",
                                    help="Get a free key at https://aistudio.google.com/apikey")
        else:
            st.success("API key loaded")
        st.caption(f"Model: `{get_model_name()}`")
        st.markdown("---")
        st.caption("Your resume is sent to Google's Gemini API for analysis and is not stored by this app.")

    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    job_desc = st.text_area("Job description (optional, improves keyword matching)", height=150)

    if st.button("Analyse resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            return
        data = uploaded.getvalue()
        if len(data) > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is larger than {MAX_FILE_MB} MB.")
            return
        try:
            text = extract_text(uploaded.name, data)
        except Exception as e:
            st.error(f"Could not read the file: {e}")
            return
        if len(text) < 100:
            st.error("Very little text was found. If this is a scanned/image PDF, "
                     "an ATS cannot read it either, so export a text-based PDF or use DOCX.")
            return
        try:
            with st.spinner("Analysing with Gemini..."):
                st.session_state["result"] = analyse_resume(text, job_desc.strip(), api_key, get_model_name())
        except json.JSONDecodeError:
            st.error("The AI returned an unreadable response. Please try again.")
            return
        except Exception as e:
            st.error(f"Gemini request failed: {e}")
            return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
